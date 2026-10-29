import hashlib
import io
import json
import os
import socket
import subprocess
import sys
import zipfile
from pathlib import Path

import pytest
import tiktoken
from docx import Document
from PIL import Image, ImageDraw
from reportlab.pdfgen.canvas import Canvas

from strahov_rag_ingestion.__main__ import execute
from strahov_rag_ingestion.chunking import Chunker, embedding_input
from strahov_rag_ingestion.contracts import ExtractionError
from strahov_rag_ingestion.io_utils import write_exclusive
from strahov_rag_ingestion.robots import evaluate
from strahov_rag_ingestion.safety import inspect_input
from strahov_rag_ingestion.worker import plain_block, run


def pdf(pages=1, columns=False, image=False):
    stream = io.BytesIO()
    canvas = Canvas(stream, pagesize=(595, 842))
    for index in range(pages):
        if image:
            from reportlab.lib.utils import ImageReader
            picture = Image.new("RGB", (400, 600), "white")
            ImageDraw.Draw(picture).text((20, 20), "SCAN 123 NO", fill="black")
            canvas.drawImage(ImageReader(picture), 30, 30, 400, 600)
        else:
            for line in range(5):
                canvas.drawString(30, 780 - line * 20, f"PAGE{index + 1} LEFT{line}: amount 1200.50; not paid.")
                if columns:
                    canvas.drawString(320, 780 - line * 20, f"RIGHT{line}: 300 CZK")
        canvas.showPage()
    canvas.save()
    return stream.getvalue()


def docx_bytes(rows=3):
    stream = io.BytesIO()
    doc = Document()
    doc.add_heading("Pravidla", 1)
    doc.add_paragraph("Před tabulkou: sleva neplatí, částka 123,50 Kč.")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Položka"
    table.rows[0].cells[1].text = "Cena (Kč)"
    for i in range(rows):
        cells = table.add_row().cells
        cells[0].text = f"Řádek {i}"
        cells[1].text = str(-1200 - i)
    doc.add_paragraph("Za tabulkou: nezaplatit podruhé.")
    doc.save(stream)
    return stream.getvalue()


def test_plain_unicode_complete_budget_and_stability(request_for):
    text = 'Sleva neplatí. Částka 1 234,50 Kč. 😀 "citace"\n' * 30
    request = request_for(text.encode(), limits={"maxChunkTokens": 70, "maxEmbeddingInputBytes": 400})
    result = run(request)
    assert result["status"] == "ok", result.get("errorCode")
    assert "".join(c["text"] for c in result["chunks"]) == text
    encoder = tiktoken.get_encoding("cl100k_base")
    for chunk in result["chunks"]:
        value = embedding_input(request.title, chunk["headingPath"], chunk["text"])
        assert chunk["tokenCount"] == len(encoder.encode(value, disallowed_special=())) <= 70
        assert len(value.encode()) <= 400
        assert chunk["blockIds"] == [result["blocks"][0]["id"]]
    blocks = [plain_block(request, "text", "Nemění se", "page 1"), plain_block(request, "text", "Nemění se", "page 2", 1)]
    before = Chunker(request).chunks(blocks)
    after = Chunker(request).chunks([plain_block(request, "text", "Nový odstavec", "page 1"), *blocks])
    assert [c["id"] for c in before] == [c["id"] for c in after[1:]]
    assert len({c["id"] for c in before}) == 2


def test_full_native_pdf_and_page_limit(request_for):
    request = request_for(pdf(3), "pages.pdf")
    result = run(request)
    assert result["status"] == "ok", result
    assert result["quality"]["processedPages"] == 3
    assert {span["page"] for block in result["blocks"] for span in block["spans"]} == {1, 2, 3}
    assert all(f"PAGE{i}" in " ".join(b["text"] for b in result["blocks"]) for i in (1,2,3))
    request.limits.maxPages = 2
    limited = run(request)
    assert limited["status"] != "ok" and not limited["document"]["complete"]
    assert limited["errorCode"] == "page-limit"


@pytest.mark.parametrize("kwargs,issue", [({"columns": True}, "pdf-needs-layout"), ({"image": True}, "pdf-page-needs-ocr")])
def test_ambiguous_pdf_requires_real_models(request_for, kwargs, issue):
    result = run(request_for(pdf(**kwargs), "complex.pdf"))
    assert result["status"] == "review-required", result
    assert issue in result["quality"]["issues"]
    assert "standard-model-profile-unavailable" in result["quality"]["issues"]
    assert not result["document"]["complete"]


def test_docx_order_tables_negations_and_no_fake_pages(request_for):
    result = run(request_for(docx_bytes(), "tables.docx"))
    assert result["status"] == "ok", result
    assert result["quality"]["tableCount"] == 1
    blocks = result["blocks"]
    before = next(i for i,b in enumerate(blocks) if "Před tabulkou" in b["text"])
    table = next(i for i,b in enumerate(blocks) if b["kind"] == "table")
    after = next(i for i,b in enumerate(blocks) if "Za tabulkou" in b["text"])
    assert before < table < after
    assert "sleva neplatí" in blocks[before]["text"]
    assert "-1202" in blocks[table]["text"]
    assert result["document"]["pageCount"] is None
    assert all("page" not in span for block in blocks for span in block["spans"])


def test_long_table_repeats_headers_and_keeps_rows(request_for):
    result = run(request_for(docx_bytes(30), "long.docx", limits={"maxChunkTokens": 140}))
    assert result["status"] == "ok", result
    table = next(b for b in result["blocks"] if b["kind"] == "table")
    chunks = [c for c in result["chunks"] if table["id"] in c["blockIds"]]
    assert len(chunks) > 1
    assert all("Cena (Kč)" in c["text"] for c in chunks)
    for row in range(30):
        assert sum(f'"Řádek {row}"' in c["text"] for c in chunks) == 1


def test_impossible_table_budget_is_explicit(request_for):
    result = run(request_for(docx_bytes(), "small.docx", limits={"maxChunkTokens": 10}))
    assert result["status"] == "review-required"
    assert result["errorCode"] in {"table-header-exceeds-budget", "table-row-exceeds-budget", "embedding-context-exceeds-budget"}


def test_html_charset_navigation_links_and_no_fetch(request_for, monkeypatch):
    def forbidden(*args, **kwargs):
        raise AssertionError("unexpected network")
    monkeypatch.setattr(socket.socket, "connect", forbidden)
    data = '<html><head><meta charset="windows-1250"></head><body><nav><a href="/trash">Trash</a></nav><main><h1>Dokumenty</h1><p>Sleva neplatí: 250 Kč.</p><a href="/rules.pdf">Řád</a></main></body></html>'.encode("cp1250")
    result = run(request_for(data, "web.html"))
    assert result["status"] == "ok", result
    assert "neplatí" in " ".join(b["text"] for b in result["blocks"])
    assert [link["href"] for link in result["links"]] == ["https://example.test/rules.pdf"]
    assert "Trash" not in " ".join(b["text"] for b in result["blocks"])


def test_feed_deduplicates_and_xml_entities_rejected(request_for):
    feed = b'<rss><channel><item><title>A</title><link>https://example.test/a</link></item><item><title>A</title><link>https://example.test/a</link></item></channel></rss>'
    result = run(request_for(feed, "feed.xml"))
    assert result["status"] == "ok" and len(result["links"]) == 1
    malicious = b'<?xml version="1.0"?><!DOCTYPE rss [<!ENTITY x SYSTEM "file:///secret">]><rss>&x;</rss>'
    result = run(request_for(malicious, "evil.xml"))
    assert result["status"] == "failed" and result["errorCode"] == "unsafe-or-invalid-xml"


@pytest.mark.parametrize("name", ["../outside.xml", "/absolute.xml", "C:/drive.xml"])
def test_zip_traversal_rejected(request_for, name):
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as z:
        z.writestr(name, "<root/>")
    result = run(request_for(buffer.getvalue(), "evil.docx"))
    assert result["errorCode"] == "unsafe-zip-path"


def test_zip_expansion_and_signature_guard(request_for):
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("large.xml", " " * 10000)
    result = run(request_for(buffer.getvalue(), "evil.docx", limits={"maxExpandedBytes": 100}))
    assert result["errorCode"] == "expanded-size-limit"
    result = run(request_for(b"plain disguised", "fake.docx"))
    assert result["status"] == "unsupported"


def test_tokenizer_missing_fails_without_fetch(request_for, monkeypatch):
    monkeypatch.delenv("TIKTOKEN_CACHE_DIR", raising=False)
    result = run(request_for(b"No network"))
    assert result["errorCode"] == "tokenizer-unavailable"


def test_exclusive_atomic_output(tmp_path):
    output = tmp_path / "result.json"
    write_exclusive(output, b"first")
    with pytest.raises(FileExistsError):
        write_exclusive(output, b"second")
    assert output.read_bytes() == b"first"
    assert list(tmp_path.iterdir()) == [output]


def test_cli_process_and_cleanup(request_for):
    request = request_for("Sleva neplatí 123 Kč.".encode())
    output = Path(request.workDir) / "result.json"
    result = execute(request, output)
    assert result["status"] == "ok", result
    assert json.loads(output.read_text(encoding="utf-8"))["status"] == "ok"
    assert not (Path(request.workDir) / ".worker-claim").exists()
    assert not (Path(request.workDir) / ".worker-result.json").exists()
    with pytest.raises(ValueError):
        execute(request, output)


def test_robots_longest_wildcard_anchor_group_and_lazy_import(tmp_path):
    body = b'User-agent: *\nDisallow: /\nUser-agent: Strahov\nDisallow: /private/*\nAllow: /private/public\nDisallow: /exact$\nCrawl-delay: 1.5\n'
    path = tmp_path / "robots.txt"
    path.write_bytes(body)
    urls = ["https://example.test/private/a", "https://example.test/private/public/x", "https://example.test/exact", "https://example.test/exact/more"]
    result = evaluate(dict(schemaVersion=1, bodyPath=str(path), bodySha256=hashlib.sha256(body).hexdigest(),userAgent="Strahov/1.0", urls=urls))
    assert [x["allowed"] for x in result["decisions"]] == [False, True, False, True]
    assert all(x["crawlDelaySeconds"] == 1.5 for x in result["decisions"])


def test_markdown_csv_tsv(request_for):
    for name, text in [("test.md", "# Nadpis\n\nSleva **neplatí**, cena 250 Kč."),
                       ("test.csv", 'Položka,Cena\n"A, B",-250\nC,125'),
                       ("test.tsv", 'Položka\tCena\nA\t-250\nB\t125')]:
        result = run(request_for(text.encode(), name))
        assert result["status"] == "ok", result
        assert result["quality"]["outputTextCharacters"] > 10
        if name.endswith(("csv", "tsv")):
            assert result["quality"]["tableCount"] == 1
            assert "-250" in result["blocks"][0]["text"]


def test_xlsx_all_sheets_and_pptx_all_slides(request_for):
    from openpyxl import Workbook
    from pptx import Presentation
    workbook = Workbook()
    workbook.active.title = "První"
    workbook.active.append(["Položka", "Kč"])
    workbook.active.append(["Sleva neplatí", -250])
    second = workbook.create_sheet("Druhý")
    second.append(["Druhý list", 125])
    data = io.BytesIO()
    workbook.save(data)
    result = run(request_for(data.getvalue(), "sheets.xlsx"))
    assert result["status"] == "ok", result
    assert result["quality"]["processedPages"] == 2
    assert "Druhý list" in " ".join(b["text"] for b in result["blocks"])
    presentation = Presentation()
    for i in range(2):
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = f"Snímek {i + 1}"
        slide.placeholders[1].text = f"Sleva neplatí {i + 100} Kč."
    data = io.BytesIO()
    presentation.save(data)
    result = run(request_for(data.getvalue(), "slides.pptx"))
    assert result["status"] == "ok", result
    assert result["quality"]["processedPages"] == 2
    assert "Snímek 2" in " ".join(b["text"] for b in result["blocks"])


def test_opendocument_text_and_spreadsheet(request_for):
    from odfdo import Document as OdfDocument, Paragraph, Table, Row, Cell
    document = OdfDocument("text")
    document.body.append(Paragraph("Sleva neplatí: 135 Kč."))
    data = io.BytesIO()
    document.save(data)
    result = run(request_for(data.getvalue(), "text.odt"))
    assert result["status"] == "ok", result
    assert "neplatí" in " ".join(b["text"] for b in result["blocks"])
    document = OdfDocument("spreadsheet")
    for name in ["První", "Druhý"]:
        table = Table(name)
        row = Row()
        row.append(Cell(name))
        row.append(Cell(-135))
        table.append(row)
        document.body.append(table)
    data = io.BytesIO()
    document.save(data)
    result = run(request_for(data.getvalue(), "sheets.ods"))
    assert result["status"] == "ok", result
    # odfdo's spreadsheet template includes its initial empty Sheet1.
    assert result["quality"]["processedPages"] == 3


def test_tracked_changes_and_legacy_explicit(request_for):
    source = io.BytesIO(docx_bytes())
    target = io.BytesIO()
    with zipfile.ZipFile(source) as original, zipfile.ZipFile(target, "w") as changed:
        for name in original.namelist():
            content = original.read(name)
            if name == "word/document.xml":
                content = content.replace(b"</w:body>", b'<w:ins><w:r><w:t>Neprijata zmena</w:t></w:r></w:ins></w:body>')
            changed.writestr(name, content)
    result = run(request_for(target.getvalue(), "tracked.docx"))
    assert result["status"] == "review-required"
    assert "unaccepted-tracked-changes" in result["quality"]["issues"]
    result = run(request_for(b"{\\rtf1 Hello}", "old.rtf"))
    assert result["status"] == "unsupported"
    assert result["errorCode"] == "legacy-converter-unavailable"


def test_timeout_and_output_limit(request_for):
    request = request_for(pdf(3), "slow.pdf", limits={"timeoutSeconds": 1})
    result = execute(request, Path(request.workDir) / "result.json")
    assert result["errorCode"] == "worker-timeout"
    assert not (Path(request.workDir) / ".worker-claim").exists()
    request = request_for(("Content " * 1000).encode(), "large.txt", limits={"maxOutputBytes": 2048})
    result = run(request)
    assert result["status"] == "failed" and result["errorCode"] == "output-size-limit"


def test_multiple_html_articles_and_vertical_table_context(request_for):
    html = b'<body><article><p>Article A</p><a href="/a.pdf">A</a></article><article><p>Article B</p><a href="/b.pdf">B</a></article></body>'
    result = run(request_for(html, "articles.html"))
    assert result["status"] == "ok", result
    assert "Article B" in " ".join(b["text"] for b in result["blocks"])
    assert len(result["links"]) == 2
    doc = Document()
    table = doc.add_table(rows=20, cols=3)
    for cell, text in zip(table.rows[0].cells, ["Kategorie", "Položka", "Kč"]):
        cell.text = text
    table.cell(1, 0).merge(table.cell(19, 0)).text = "Student"
    for i in range(1, 20):
        table.cell(i, 1).text = f"Cena {i}"
        table.cell(i, 2).text = str(100 + i)
    stream = io.BytesIO()
    doc.save(stream)
    result = run(request_for(stream.getvalue(), "merged.docx", limits={"maxChunkTokens": 190}))
    assert result["status"] == "ok", result
    assert len(result["chunks"]) > 1
    assert all("Student" in chunk["text"] for chunk in result["chunks"])


def test_feed_limit_and_owned_cleanup(request_for):
    feed = '<rss><channel>' + ''.join(f'<item><title>{i}</title><link>https://example.test/document/{i}</link></item>' for i in range(200)) + '</channel></rss>'
    request = request_for(feed.encode(), "largefeed.xml", limits={"maxOutputBytes": 2048})
    output = Path(request.workDir) / "result.json"
    result = execute(request, output)
    assert result["errorCode"] == "output-size-limit", result
    assert output.stat().st_size <= 2048
    request = request_for(b"sentinel", "sentinel.txt")
    sentinel = Path(request.workDir) / ".worker-result.json"
    sentinel.write_bytes(b"belongs to previous job")
    with pytest.raises(ValueError):
        execute(request, Path(request.workDir) / "result.json")
    assert sentinel.read_bytes() == b"belongs to previous job"


def test_actual_module_entrypoint_and_robots_lazy_import(request_for):
    request = request_for(docx_bytes(), "cli.docx")
    path = Path(request.workDir) / "request.json"
    output = Path(request.workDir) / "result.json"
    path.write_text(request.model_dump_json(), encoding="utf-8")
    process = subprocess.run([sys.executable, "-m", "strahov_rag_ingestion", "--request", str(path), "--output", str(output)], capture_output=True, text=True, timeout=70)
    assert process.returncode == 0, (process.stdout, process.stderr)
    assert json.loads(output.read_text(encoding="utf-8"))["quality"]["tableCount"] == 1
    assert "neplatí" not in process.stdout + process.stderr
    lazy = subprocess.run([sys.executable, "-c", 'import sys; import strahov_rag_ingestion.robots; assert not any(x.startswith(("docling", "torch", "tiktoken")) for x in sys.modules)'], capture_output=True, timeout=10)
    assert lazy.returncode == 0, lazy.stderr


def test_opendocument_slides(request_for):
    from odfdo import Document as OdfDocument, DrawPage, Frame
    document = OdfDocument("presentation")
    for i in range(2):
        page = DrawPage(name=f"Slide{i + 1}")
        frame = Frame(name=f"Frame{i + 1}", size=("10cm", "5cm"))
        frame.set_text_box(f"Snímek {i + 1}: sleva neplatí 250 Kč.")
        page.append(frame)
        document.body.append(page)
    stream = io.BytesIO()
    document.save(stream)
    result = run(request_for(stream.getvalue(), "slides.odp"))
    assert result["status"] == "ok", result
    assert "Snímek 2" in " ".join(b["text"] for b in result["blocks"])


def test_formula_cache_and_failure_envelope(request_for):
    from odfdo import Document as OdfDocument, Table, Row, Cell
    from pydantic import ValidationError
    doc = OdfDocument("spreadsheet")
    table, row = Table("Formula"), Row()
    row.append(Cell(999, formula="of:=1+1"))
    table.append(row)
    doc.body.append(table)
    stream = io.BytesIO()
    doc.save(stream)
    result = run(request_for(stream.getvalue(), "formula.ods"))
    assert result["status"] == "review-required"
    assert "spreadsheet-formula-values-require-review" in result["quality"]["issues"]
    with pytest.raises(ValidationError):
        request_for(b"text", "small-envelope.txt", title="A" * 2000, limits={"maxOutputBytes": 2048})


def test_picture_children_list_markers_and_missing_model(request_for):
    from docling_core.types.doc import DoclingDocument, DocItemLabel, GroupLabel
    from strahov_rag_ingestion.docling_adapter import normalize_doc
    document = DoclingDocument(name="synthetic")
    picture = document.add_picture()
    document.add_text(label=DocItemLabel.TEXT, text="Text uvnitř obrázku 123 Kč.", parent=picture)
    list_group = document.add_group(label=GroupLabel.LIST)
    document.add_list_item(text="Sleva neplatí.", enumerated=True, marker="(7)", parent=list_group)
    blocks, issues, _ = normalize_doc(document, "pdf", request_for(b"placeholder"))
    assert any(b["text"] == "(7) Sleva neplatí." for b in blocks)
    assert any("Text uvnitř obrázku" in b["text"] for b in blocks)
    assert "unrepresented-image" in issues
    result = run(request_for(pdf(), "missing.pdf", mode="standard"))
    assert result["status"] == "review-required"
    assert result["errorCode"] == "standard-model-profile-unavailable"
